"""
Google Drive Upload Module
- Upload file Excel lên Google Drive
- Convert thành Google Sheet format
- Quản lý folder trên Drive
"""

import os
import time
from typing import Optional
from pathlib import Path

from config import EXPORTS_DIR


class GoogleDriveUploader:
    """Upload files to Google Drive and convert to Google Sheets format."""

    def __init__(self, log_callback=None):
        self.log = log_callback or print
        self._service = None

    def connect(self) -> bool:
        """Establish connection to Google Drive API."""
        try:
            from googleapiclient.discovery import build
            from cloud.google_sheets import _get_google_credentials
            creds = _get_google_credentials()
            self._service = build('drive', 'v3', credentials=creds)
            self.log("[OK] Kết nối Google Drive API thành công")
            return True
        except Exception as e:
            self.log(f"[ERROR] Kết nối Drive thất bại: {e}")
            return False

    def upload_excel_as_sheet(self, file_path: str,
                               folder_id: str = None,
                               filename: str = None) -> Optional[str]:
        """Upload Excel file to Drive and convert to Google Sheets.

        Returns the Google Drive file ID.
        """
        if not self._service:
            self.log("[ERROR] Chưa kết nối Drive")
            return None

        file_path = Path(file_path)
        if not file_path.exists():
            self.log(f"[ERROR] File không tồn tại: {file_path}")
            return None

        try:
            from googleapiclient.http import MediaFileUpload

            file_metadata = {
                'name': filename or file_path.stem,
                'mimeType': 'application/vnd.google-apps.spreadsheet',
            }
            if folder_id:
                file_metadata['parents'] = [folder_id]

            media = MediaFileUpload(
                str(file_path),
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                resumable=True
            )

            file = self._service.files().create(
                body=file_metadata,
                media_body=media,
                fields='id, webViewLink'
            ).execute()

            file_id = file.get('id')
            web_link = file.get('webViewLink', '')

            self.log(f"[OK] Đã upload lên Drive: {file_path.name}")
            self.log(f"  [INFO] Link: {web_link}")
            return file_id

        except Exception as e:
            self.log(f"[ERROR] Upload Drive thất bại: {e}")
            return None

    def upload_file(self, file_path: str, folder_id: str = None,
                     filename: str = None) -> Optional[str]:
        """Upload any file to Drive (without conversion)."""
        if not self._service:
            return None

        file_path = Path(file_path)
        if not file_path.exists():
            return None

        try:
            from googleapiclient.http import MediaFileUpload

            mime_types = {
                '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                '.xls': 'application/vnd.ms-excel',
                '.csv': 'text/csv',
                '.pdf': 'application/pdf',
                '.json': 'application/json',
            }
            mime = mime_types.get(file_path.suffix.lower(), 'application/octet-stream')

            file_metadata = {'name': filename or file_path.name}
            if folder_id:
                file_metadata['parents'] = [folder_id]

            media = MediaFileUpload(str(file_path), mimetype=mime, resumable=True)

            file = self._service.files().create(
                body=file_metadata,
                media_body=media,
                fields='id'
            ).execute()

            self.log(f"[OK] Uploaded: {file_path.name}")
            return file.get('id')

        except Exception as e:
            self.log(f"[ERROR] Upload failed: {e}")
            return None

    def create_folder(self, folder_name: str,
                       parent_id: str = None) -> Optional[str]:
        """Create a folder on Drive. Returns folder ID."""
        if not self._service:
            return None

        try:
            metadata = {
                'name': folder_name,
                'mimeType': 'application/vnd.google-apps.folder',
            }
            if parent_id:
                metadata['parents'] = [parent_id]

            folder = self._service.files().create(
                body=metadata, fields='id'
            ).execute()

            folder_id = folder.get('id')
            self.log(f"[OK] Tạo folder Drive: {folder_name}")
            return folder_id

        except Exception as e:
            self.log(f"[ERROR] Create folder: {e}")
            return None

    def list_files(self, folder_id: str = None, max_results: int = 50) -> list:
        """List files in a Drive folder."""
        if not self._service:
            return []

        try:
            query = ""
            if folder_id:
                query = f"'{folder_id}' in parents and trashed = false"
            else:
                query = "trashed = false"

            results = self._service.files().list(
                q=query,
                pageSize=max_results,
                fields="files(id, name, mimeType, modifiedTime, webViewLink)"
            ).execute()

            return results.get('files', [])

        except Exception as e:
            self.log(f"[ERROR] List files: {e}")
            return []


def export_data_to_excel(data: list[dict], output_path: str) -> bool:
    """Export data to Excel file (.xlsx) using openpyxl."""
    try:
        from openpyxl import Workbook
        from config import GRID_COLUMNS

        wb = Workbook()
        ws = wb.active
        ws.title = "Gmail Data"

        # Header
        headers = [col["label"] for col in GRID_COLUMNS]
        ws.append(headers)

        # Style header
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2E7D32", fill_type="solid")

        # Data rows
        from config import GRID_COLUMN_KEYS
        for record in data:
            row = [str(record.get(key, "")) for key in GRID_COLUMN_KEYS]
            ws.append(row)

        # Auto-width
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

        wb.save(output_path)
        return True

    except Exception as e:
        print(f"[ERROR] Export Excel: {e}")
        return False

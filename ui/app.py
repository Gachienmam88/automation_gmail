"""
Main Application Window
- Desktop app with Grid UI
- Control panel: Start/Pause/Stop
- Log panel
- Settings dialog
- Cloud sync integration
"""

import os
import sys
import time
import threading
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, scrolledtext
from pathlib import Path

# Add project root to path
sys.path.insert(0, str(Path(__file__).parent.parent))

from config import (
    GRID_COLUMN_KEYS, ACTIONS, DISPLAY_MODES,
    MAX_THREADS, PROFILES_DIR, EXPORTS_DIR,
    load_settings, save_settings, apply_settings,
)
from ui.grid_view import GridView
from ui.dialogs import SettingsDialog, SheetSyncDialog
from profile.firefox_manager import FirefoxProfileManager
from security.two_factor import TwoFactorManager
from core.engine import AutomationEngine
from core.gmail_actions import execute_action
from cloud.google_sheets import GoogleSheetsSync
from cloud.google_drive import GoogleDriveUploader, export_data_to_excel


class GmailAutomationApp:
    """Main application class."""

    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Gmail Automation Tool - TechLaAI")
        self.root.geometry("1400x800")
        self.root.minsize(1000, 600)

        # Set icon if available
        try:
            icon_path = Path(__file__).parent.parent / "icon.ico"
            if icon_path.exists():
                self.root.iconbitmap(str(icon_path))
        except Exception:
            pass

        # ── Core modules ──
        self.profile_manager = FirefoxProfileManager(
            log_callback=self._log
        )
        self.twofa_manager = TwoFactorManager(log_callback=self._log)
        self.engine = AutomationEngine(
            log_callback=self._log,
            progress_callback=self._on_task_progress,
        )
        self.sheets_sync = GoogleSheetsSync(log_callback=self._log)
        self.drive_uploader = GoogleDriveUploader(log_callback=self._log)

        # State
        self._is_running = False
        self._task_index_map = {}  # gmail → row index

        # Load settings
        settings = load_settings()
        apply_settings(settings)

        # Build UI
        self._build_ui()
        self._apply_theme()

        # Load saved data if any
        self._load_saved_data()

    def _build_ui(self):
        """Build the main application UI."""
        # ── Menu Bar ──
        menubar = tk.Menu(self.root)
        self.root.config(menu=menubar)

        file_menu = tk.Menu(menubar, tearoff=0)
        file_menu.add_command(label="Import from File...",
                               command=self._import_from_file)
        file_menu.add_command(label="Export to Excel...",
                               command=self._export_to_excel)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=self._on_close)
        menubar.add_cascade(label="File", menu=file_menu)

        cloud_menu = tk.Menu(menubar, tearoff=0)
        cloud_menu.add_command(label="Google Sheet Sync...",
                                command=self._open_sheet_sync)
        cloud_menu.add_command(label="Upload to Drive...",
                                command=self._upload_to_drive)
        menubar.add_cascade(label="Cloud", menu=cloud_menu)

        tools_menu = tk.Menu(menubar, tearoff=0)
        tools_menu.add_command(label="Manage Profiles",
                                command=self._manage_profiles)
        tools_menu.add_command(label="Check IP",
                                command=self._check_ip)
        tools_menu.add_command(label="Test 2FA",
                                command=self._test_2fa)
        menubar.add_cascade(label="Tools", menu=tools_menu)

        settings_menu = tk.Menu(menubar, tearoff=0)
        settings_menu.add_command(label="Settings...",
                                   command=self._open_settings)
        menubar.add_cascade(label="Settings", menu=settings_menu)

        # ── Main Paned Window ──
        paned = ttk.PanedWindow(self.root, orient=tk.VERTICAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        # ── Top: Grid + Controls ──
        top_frame = ttk.Frame(paned)
        paned.add(top_frame, weight=3)

        # Control Panel
        control_frame = ttk.LabelFrame(top_frame, text="Control Panel")
        control_frame.pack(fill=tk.X, padx=5, pady=(0, 5))

        # Row 1: Main controls
        ctrl_row1 = ttk.Frame(control_frame)
        ctrl_row1.pack(fill=tk.X, padx=5, pady=3)

        self.btn_start = ttk.Button(
            ctrl_row1, text="▶ START", command=self._start_automation
        )
        self.btn_start.pack(side=tk.LEFT, padx=5)

        self.btn_pause = ttk.Button(
            ctrl_row1, text="⏸ PAUSE", command=self._pause_automation,
            state=tk.DISABLED
        )
        self.btn_pause.pack(side=tk.LEFT, padx=5)

        self.btn_stop = ttk.Button(
            ctrl_row1, text="⏹ STOP", command=self._stop_automation,
            state=tk.DISABLED
        )
        self.btn_stop.pack(side=tk.LEFT, padx=5)

        ttk.Separator(ctrl_row1, orient=tk.VERTICAL).pack(
            side=tk.LEFT, fill=tk.Y, padx=10)

        # Thread count
        ttk.Label(ctrl_row1, text="Threads:").pack(side=tk.LEFT)
        self.thread_spin = ttk.Spinbox(
            ctrl_row1, from_=1, to=10, width=5
        )
        self.thread_spin.set(MAX_THREADS)
        self.thread_spin.pack(side=tk.LEFT, padx=5)

        # Progress
        self.progress_var = tk.StringVar(value="Ready")
        ttk.Label(ctrl_row1, textvariable=self.progress_var).pack(
            side=tk.RIGHT, padx=10)

        self.progress_bar = ttk.Progressbar(
            ctrl_row1, mode="determinate", length=200
        )
        self.progress_bar.pack(side=tk.RIGHT, padx=5)

        # Grid
        self.grid_view = GridView(
            top_frame, on_data_change=self._on_data_change
        )
        self.grid_view.pack(fill=tk.BOTH, expand=True)

        # ── Bottom: Log Panel ──
        log_frame = ttk.LabelFrame(paned, text="Log")
        paned.add(log_frame, weight=1)

        self.log_text = scrolledtext.ScrolledText(
            log_frame, height=10, wrap=tk.WORD,
            font=("Consolas", 9),
            bg="#1E1E1E", fg="#D4D4D4",
            insertbackground="#D4D4D4",
        )
        self.log_text.pack(fill=tk.BOTH, expand=True, padx=3, pady=3)

        # Log tags
        self.log_text.tag_configure("INFO", foreground="#4FC3F7")
        self.log_text.tag_configure("OK", foreground="#81C784")
        self.log_text.tag_configure("WARNING", foreground="#FFB74D")
        self.log_text.tag_configure("ERROR", foreground="#E57373")
        self.log_text.tag_configure("DEBUG", foreground="#90A4AE")

        # ── Status Bar ──
        status_bar = ttk.Frame(self.root)
        status_bar.pack(fill=tk.X, padx=5, pady=(0, 5))

        self.status_label = ttk.Label(
            status_bar, text="Gmail Automation Tool v1.0 - TechLaAI"
        )
        self.status_label.pack(side=tk.LEFT)

        self.ip_label = ttk.Label(status_bar, text="IP: --")
        self.ip_label.pack(side=tk.RIGHT, padx=10)

        # Close handler
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

    def _apply_theme(self):
        """Apply a modern theme."""
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass

        # Custom styles
        style.configure("Treeview",
                         rowheight=25,
                         font=("Segoe UI", 9))
        style.configure("Treeview.Heading",
                         font=("Segoe UI", 9, "bold"))

    # ─── Logging ───

    def _log(self, message: str):
        """Thread-safe logging to the text widget."""
        def _append():
            self.log_text.insert(tk.END, message + "\n")

            # Apply tag based on prefix
            line_start = self.log_text.index("end-2l linestart")
            line_end = self.log_text.index("end-2l lineend")
            for tag in ["OK", "ERROR", "WARNING", "INFO", "DEBUG"]:
                if f"[{tag}]" in message:
                    self.log_text.tag_add(tag, line_start, line_end)
                    break

            self.log_text.see(tk.END)

        self.root.after(0, _append)

    # ─── Automation Control ───

    def _start_automation(self):
        """Start automation for selected rows."""
        selected_data = self.grid_view.get_selected_data()
        selected_indices = self.grid_view.get_selected_indices()

        if not selected_data:
            # If nothing selected, use all rows
            selected_data = self.grid_view.get_data()
            selected_indices = list(range(len(selected_data)))

        if not selected_data:
            messagebox.showwarning("Warning", "Không có dữ liệu để xử lý!")
            return

        # Validate data
        valid_tasks = []
        for i, (data, idx) in enumerate(zip(selected_data, selected_indices)):
            if not data.get("gmail"):
                continue
            if not data.get("password"):
                self._log(f"[WARNING] {data['gmail']}: thiếu password, bỏ qua")
                continue

            # Prepare task
            task = dict(data)
            task["action"] = task.get("action") or self.grid_view.action_combo.get()
            task["headless"] = (
                task.get("display_mode", "") == "Headless"
                or self.grid_view.display_combo.get() == "Headless"
            )

            # Get/create profile
            profile_path = self.profile_manager.get_or_create_profile(
                task["gmail"]
            )
            if profile_path:
                task["profile_path"] = profile_path

            self._task_index_map[task["gmail"]] = idx
            valid_tasks.append(task)

        if not valid_tasks:
            messagebox.showwarning("Warning", "Không có task hợp lệ!")
            return

        self._is_running = True
        self.btn_start.config(state=tk.DISABLED)
        self.btn_pause.config(state=tk.NORMAL)
        self.btn_stop.config(state=tk.NORMAL)
        self.progress_bar["maximum"] = len(valid_tasks)
        self.progress_bar["value"] = 0
        self.progress_var.set(f"Running: 0/{len(valid_tasks)}")

        max_threads = int(self.thread_spin.get())

        self._log(f"[INFO] Bắt đầu {len(valid_tasks)} tasks "
                  f"({max_threads} threads)")

        # Start in background thread
        def run():
            self.engine.max_threads = max_threads
            self.engine.start_batch(valid_tasks, self._execute_single_task)

            # Wait for completion
            while self.engine.is_running():
                time.sleep(1)

            self.root.after(0, self._on_automation_complete)

        threading.Thread(target=run, daemon=True).start()

    def _execute_single_task(self, browser, task, log, result):
        """Execute a single task (proxy is already configured in BrowserInstance)."""
        try:
            execute_action(browser, task, log, result)
        except Exception as e:
            result.error = str(e)
            log(f"[ERROR] Task execution failed: {e}")

    def _on_task_progress(self, gmail: str, status: str, result, **kwargs):
        """Callback when a task's status changes."""
        try:
            idx = self._task_index_map.get(gmail)
        except Exception:
            return
        if idx is None:
            return

        status_text = {
            "running": "Running...",
            "done": "OK",
            "error": f"Error: {result.error if result else ''}",
        }.get(status, status)

        result_text = ""
        if result:
            result_text = result.message or result.error or ""

        def update_ui():
            self.grid_view.set_row_status(idx, status_text, result_text)

            # Update result columns if applicable
            if result and result.success:
                updates = {"status": status_text, "result": result_text}
                if result.new_password:
                    updates["password"] = result.new_password
                    updates["new_password"] = ""
                if result.twofa_secret:
                    updates["twofa_secret"] = result.twofa_secret
                if result.timestamp:
                    updates["last_login"] = result.timestamp
                self.grid_view.update_row(idx, updates)

            # Update progress
            completed = sum(
                1 for d in self.grid_view.get_data()
                if d.get("status", "").lower() in ("ok", "error", "done")
                or "error" in d.get("status", "").lower()
            )
            total = len(self._task_index_map)
            self.progress_bar["value"] = completed
            self.progress_var.set(f"Running: {completed}/{total}")

            # Update Sheet status
            if self.sheets_sync.sheet_id:
                try:
                    self.sheets_sync.update_row_status(
                        idx, status_text, result_text
                    )
                except Exception:
                    pass

        self.root.after(0, update_ui)

    def _on_automation_complete(self):
        """Called when all tasks are done."""
        if not self._is_running:
            return  # Already completed, avoid duplicate calls
        self._is_running = False
        self.btn_start.config(state=tk.NORMAL)
        self.btn_pause.config(state=tk.DISABLED)
        self.btn_stop.config(state=tk.DISABLED)
        self.progress_var.set("Completed")
        self._log("[DONE] Tất cả tasks đã hoàn thành!")

        # Auto export if enabled
        self._auto_export_results()

    def _pause_automation(self):
        """Pause/Resume automation."""
        if self.btn_pause.cget("text") == "⏸ PAUSE":
            self.engine.pause()
            self.btn_pause.config(text="▶ RESUME")
            self.progress_var.set("Paused")
        else:
            self.engine.resume()
            self.btn_pause.config(text="⏸ PAUSE")
            self.progress_var.set("Running...")

    def _stop_automation(self):
        """Stop automation."""
        if messagebox.askyesno("Xác nhận", "Dừng tất cả tasks?"):
            self.engine.stop()
            self._is_running = False
            self.btn_start.config(state=tk.NORMAL)
            self.btn_pause.config(state=tk.DISABLED, text="⏸ PAUSE")
            self.btn_stop.config(state=tk.DISABLED)
            self.progress_var.set("Stopped")

    # ─── Data Operations ───

    def _on_data_change(self):
        """Called when grid data changes."""
        self._save_data()

    def _save_data(self):
        """Save current grid data to local file."""
        import json
        data_file = Path(__file__).parent.parent / "data.json"
        try:
            with open(data_file, "w", encoding="utf-8") as f:
                json.dump(self.grid_view.get_data(), f,
                          indent=2, ensure_ascii=False)
        except Exception:
            pass

    def _load_saved_data(self):
        """Load previously saved data."""
        import json
        data_file = Path(__file__).parent.parent / "data.json"
        if data_file.exists():
            try:
                with open(data_file, "r", encoding="utf-8") as f:
                    data = json.load(f)
                self.grid_view.load_data(data)
                self._log(f"[INFO] Đã load {len(data)} dòng từ file lưu")
            except Exception:
                pass

    def _import_from_file(self):
        """Import data from Excel/CSV/TXT file."""
        file_path = filedialog.askopenfilename(
            title="Import Data",
            filetypes=[
                ("Excel files", "*.xlsx;*.xls"),
                ("CSV files", "*.csv"),
                ("Text files", "*.txt"),
                ("All files", "*.*"),
            ]
        )
        if not file_path:
            return

        try:
            data = []
            ext = os.path.splitext(file_path)[1].lower()

            if ext in ('.xlsx', '.xls'):
                from openpyxl import load_workbook
                wb = load_workbook(file_path, read_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    for i, row in enumerate(rows[1:], start=1):  # Skip header
                        record = {"stt": str(i)}
                        for j, key in enumerate(GRID_COLUMN_KEYS[1:]):
                            if j < len(row):
                                record[key] = str(row[j] or "")
                            else:
                                record[key] = ""
                        data.append(record)

            elif ext == '.csv':
                import csv
                with open(file_path, 'r', encoding='utf-8') as f:
                    reader = csv.reader(f)
                    rows = list(reader)
                    for i, row in enumerate(rows[1:], start=1):
                        record = {"stt": str(i)}
                        for j, key in enumerate(GRID_COLUMN_KEYS[1:]):
                            if j < len(row):
                                record[key] = row[j]
                            else:
                                record[key] = ""
                        data.append(record)

            elif ext == '.txt':
                with open(file_path, 'r', encoding='utf-8') as f:
                    for i, line in enumerate(f, start=1):
                        parts = line.strip().split('\t')
                        if not parts or not parts[0]:
                            continue
                        record = {"stt": str(i)}
                        # Assume: gmail\tpassword\t...
                        keys = GRID_COLUMN_KEYS[1:]
                        for j, key in enumerate(keys):
                            if j < len(parts):
                                record[key] = parts[j]
                            else:
                                record[key] = ""
                        data.append(record)

            if data:
                self.grid_view.load_data(data)
                self._log(f"[OK] Đã import {len(data)} dòng từ {os.path.basename(file_path)}")
            else:
                messagebox.showwarning("Warning", "File không có dữ liệu!")

        except Exception as e:
            messagebox.showerror("Error", f"Import thất bại: {e}")
            self._log(f"[ERROR] Import: {e}")

    def _export_to_excel(self):
        """Export grid data to Excel file."""
        file_path = filedialog.asksaveasfilename(
            title="Export to Excel",
            defaultextension=".xlsx",
            initialdir=str(EXPORTS_DIR),
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"gmail_data_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not file_path:
            return

        data = self.grid_view.get_data()
        if export_data_to_excel(data, file_path):
            self._log(f"[OK] Đã export {len(data)} dòng → {file_path}")
            messagebox.showinfo("Success", f"Đã export {len(data)} dòng!")
        else:
            messagebox.showerror("Error", "Export thất bại!")

    def _auto_export_results(self):
        """Auto export results after completion."""
        try:
            timestamp = time.strftime('%Y%m%d_%H%M%S')
            export_path = str(EXPORTS_DIR / f"result_{timestamp}.xlsx")
            data = self.grid_view.get_data()
            if export_data_to_excel(data, export_path):
                self._log(f"[OK] Auto-export: {export_path}")

                # Upload to Drive if connected
                if self.drive_uploader._service:
                    self.drive_uploader.upload_excel_as_sheet(export_path)
        except Exception as e:
            self._log(f"[WARNING] Auto-export: {e}")

    # ─── Cloud Sync ───

    def _open_sheet_sync(self):
        """Open Google Sheet sync dialog."""
        dialog = SheetSyncDialog(
            self.root,
            sheets_sync=self.sheets_sync,
            on_import=lambda data: self.grid_view.load_data(data),
            on_export=lambda: self.grid_view.get_data(),
            log_callback=self._log,
        )

    def _upload_to_drive(self):
        """Export and upload to Google Drive."""
        data = self.grid_view.get_data()
        if not data:
            messagebox.showwarning("Warning", "Không có dữ liệu!")
            return

        timestamp = time.strftime('%Y%m%d_%H%M%S')
        export_path = str(EXPORTS_DIR / f"upload_{timestamp}.xlsx")

        if not export_data_to_excel(data, export_path):
            messagebox.showerror("Error", "Export thất bại!")
            return

        def upload():
            if not self.drive_uploader._service:
                if not self.drive_uploader.connect():
                    return
            self.drive_uploader.upload_excel_as_sheet(export_path)

        threading.Thread(target=upload, daemon=True).start()

    # ─── Tools ───

    def _manage_profiles(self):
        """Show profile management dialog."""
        profiles = self.profile_manager.list_profiles()
        info = "\n".join(
            f"  {p['gmail']} - {p['size_mb']}MB"
            for p in profiles
        ) or "  (Chưa có profile)"

        messagebox.showinfo(
            "Firefox Profiles",
            f"Profiles directory: {PROFILES_DIR}\n\n{info}"
        )

    def _check_ip(self):
        """Check current IP."""
        import requests
        def check():
            try:
                resp = requests.get("https://api.ipify.org", timeout=10)
                ip = resp.text.strip()
                self.root.after(0, lambda: self.ip_label.config(
                    text=f"IP: {ip}"))
                self._log(f"[INFO] Current IP: {ip}")
            except Exception:
                self._log("[WARNING] Không thể kiểm tra IP")
        threading.Thread(target=check, daemon=True).start()

    def _test_2fa(self):
        """Test 2FA code generation."""
        selected = self.grid_view.get_selected_data()
        if not selected:
            messagebox.showinfo("Info", "Chọn dòng có 2FA Secret để test")
            return

        for row in selected:
            secret = row.get("twofa_secret", "")
            if secret:
                code = self.twofa_manager.get_totp_code(secret)
                self._log(f"[INFO] {row.get('gmail', '?')}: 2FA = {code}")

    def _open_settings(self):
        """Open settings dialog."""
        SettingsDialog(self.root, on_save=self._apply_new_settings)

    def _apply_new_settings(self, settings: dict):
        """Apply new settings from dialog."""
        apply_settings(settings)
        save_settings(settings)
        self._log("[OK] Settings đã được lưu")

    # ─── Lifecycle ───

    def _on_close(self):
        """Handle app close."""
        if self._is_running:
            if not messagebox.askyesno(
                "Xác nhận",
                "Automation đang chạy. Dừng và thoát?"
            ):
                return
            self.engine.stop()

        self._save_data()
        self.sheets_sync.stop_auto_sync()
        self.root.destroy()

    def run(self):
        """Start the application."""
        self._log("[INFO] Gmail Automation Tool khởi động")
        self._log("[INFO] Nhập dữ liệu hoặc import từ file/Google Sheet")
        self._log("[INFO] Chọn dòng → Chọn Action → Bấm START")

        # Check IP on startup
        self._check_ip()

        self.root.mainloop()
